"""Stakeholder reporting: Excel workbook and PowerPoint deck."""

import os

import matplotlib

matplotlib.use("Agg")
import matplotlib.pyplot as plt
import pandas as pd


def _charts(sentiment_summary, topic_summary, segment_profile, out_dir):
    os.makedirs(out_dir, exist_ok=True)
    paths = {}

    ax = sentiment_summary.plot(
        kind="bar", x="department", y=["pct_positive", "pct_negative"],
        figsize=(8, 4), rot=20, title="Positive vs Negative Feedback by Department",
    )
    ax.set_ylabel("share of comments")
    fig = ax.get_figure()
    fig.tight_layout()
    p = os.path.join(out_dir, "sentiment_by_department.png")
    fig.savefig(p, dpi=150)
    paths["sentiment"] = p
    plt.close(fig)

    fig, ax = plt.subplots(figsize=(8, 4))
    ax.barh(topic_summary["top_words"].str.slice(0, 32), topic_summary["n_comments"])
    ax.set_title("Comments per LDA Topic")
    ax.set_xlabel("comments")
    fig.tight_layout()
    p = os.path.join(out_dir, "topics.png")
    fig.savefig(p, dpi=150)
    paths["topics"] = p
    plt.close(fig)

    ax = segment_profile.plot(
        kind="bar", x="segment", y=["engagement", "salary_sat", "sentiment"],
        figsize=(8, 4), rot=20, title="Segment Profiles",
    )
    ax.set_ylabel("average score")
    fig = ax.get_figure()
    fig.tight_layout()
    p = os.path.join(out_dir, "segments.png")
    fig.savefig(p, dpi=150)
    paths["segments"] = p
    plt.close(fig)

    return paths


def export_excel(df, sentiment_summary, topic_summary, segment_profile, out_dir="outputs"):
    os.makedirs(out_dir, exist_ok=True)
    path = os.path.join(out_dir, "feedback_insights.xlsx")
    overall = pd.DataFrame(
        [
            {
                "n_responses": len(df),
                "mean_sentiment": df["sentiment_score"].mean().round(3),
                "pct_positive": round((df["sentiment_label"] == "positive").mean(), 3),
                "pct_negative": round((df["sentiment_label"] == "negative").mean(), 3),
            }
        ]
    )
    with pd.ExcelWriter(path, engine="openpyxl") as writer:
        overall.to_excel(writer, sheet_name="Overall", index=False)
        sentiment_summary.to_excel(writer, sheet_name="Sentiment by Dept", index=False)
        topic_summary.to_excel(writer, sheet_name="Topics", index=False)
        segment_profile.to_excel(writer, sheet_name="Segments", index=False)
    return path


def export_pptx(sentiment_summary, topic_summary, segment_profile, kmeans_meta, out_dir="outputs"):
    from pptx import Presentation
    from pptx.util import Inches

    os.makedirs(out_dir, exist_ok=True)
    path = os.path.join(out_dir, "feedback_insights.pptx")
    prs = Presentation()
    charts = _charts(sentiment_summary, topic_summary, segment_profile, out_dir)

    slide = prs.slides.add_slide(prs.slide_layouts[0])
    slide.shapes.title.text = "Employee Feedback Analytics"
    slide.placeholders[1].text = "Sentiment, Topics & Workforce Segmentation"

    slide = prs.slides.add_slide(prs.slide_layouts[5])
    slide.shapes.title.text = "Sentiment by Department"
    slide.shapes.add_picture(charts["sentiment"], Inches(1), Inches(1.5), width=Inches(8))

    slide = prs.slides.add_slide(prs.slide_layouts[5])
    slide.shapes.title.text = "Key Themes (LDA Topics)"
    slide.shapes.add_picture(charts["topics"], Inches(1), Inches(1.5), width=Inches(8))

    slide = prs.slides.add_slide(prs.slide_layouts[5])
    slide.shapes.title.text = f"Employee Segments (k={kmeans_meta['best_k']}, silhouette={kmeans_meta['silhouette']})"
    slide.shapes.add_picture(charts["segments"], Inches(1), Inches(1.5), width=Inches(8))

    slide = prs.slides.add_slide(prs.slide_layouts[1])
    slide.shapes.title.text = "Recommendations"
    bullets = slide.placeholders[1].text_frame
    lines = [
        f"Lowest sentiment department: {sentiment_summary.sort_values('mean_sentiment').iloc[0]['department']}",
        "Review workload for segments with high workload and low engagement",
        "Prioritize themes with high volume and negative sentiment",
    ]
    for i, line in enumerate(lines):
        p = bullets.paragraphs[0] if i == 0 else bullets.add_paragraph()
        p.text = line

    prs.save(path)
    return path